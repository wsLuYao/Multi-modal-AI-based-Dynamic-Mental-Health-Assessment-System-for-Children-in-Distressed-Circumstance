"""Safe, in-memory extraction for small DOCX inputs."""

from __future__ import annotations

import base64
import binascii
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from docx import Document

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
MAX_BASE64_CHARS = ((MAX_FILE_BYTES + 2) // 3) * 4
MAX_ZIP_MEMBERS = 1_024
MAX_UNCOMPRESSED_BYTES = 32 * 1024 * 1024


class DocumentValidationError(ValueError):
    """Raised when an input file fails a documented safety constraint."""


def decode_base64_document(encoded: str) -> bytes:
    if not isinstance(encoded, str) or not encoded.strip():
        raise DocumentValidationError("missing document payload")
    payload = encoded.split(",", 1)[-1].strip()
    if len(payload) > MAX_BASE64_CHARS:
        raise DocumentValidationError("encoded document exceeds the 5 MiB limit")
    try:
        raw = base64.b64decode(payload, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise DocumentValidationError("document payload is not valid base64") from exc
    if not raw:
        raise DocumentValidationError("document is empty")
    if len(raw) > MAX_FILE_BYTES:
        raise DocumentValidationError("document exceeds the 5 MiB limit")
    return raw


def extract_docx_text(raw: bytes) -> str:
    """Extract paragraphs and table cells without writing the upload to disk."""

    if not raw or len(raw) > MAX_FILE_BYTES:
        raise DocumentValidationError("document size is outside the allowed range")
    try:
        with ZipFile(BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                raise DocumentValidationError("DOCX contains too many archive members")
            names = {member.filename for member in members}
            if len(names) != len(members):
                raise DocumentValidationError("DOCX contains duplicate archive members")
            if any(member.flag_bits & 0x1 for member in members):
                raise DocumentValidationError("encrypted DOCX files are not supported")
            uncompressed_size = sum(member.file_size for member in members)
            if uncompressed_size > MAX_UNCOMPRESSED_BYTES:
                raise DocumentValidationError("expanded DOCX exceeds the 32 MiB safety limit")
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentValidationError("file is not a valid Word DOCX document")
        document = Document(BytesIO(raw))
    except DocumentValidationError:
        raise
    except (BadZipFile, KeyError, ValueError) as exc:
        raise DocumentValidationError("file is not a readable Word DOCX document") from exc

    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise DocumentValidationError("document contains no text")
    if len(text) > MAX_TEXT_CHARS:
        raise DocumentValidationError("extracted text exceeds the 100,000-character limit")
    return text
