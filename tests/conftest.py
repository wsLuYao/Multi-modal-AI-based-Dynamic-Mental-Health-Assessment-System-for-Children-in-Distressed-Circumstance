from io import BytesIO

import pytest
from docx import Document


@pytest.fixture
def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("这是完全合成的材料。孩子有时感到孤独，放学后长期独处。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "资料来源"
    table.cell(0, 1).text = "合成叙事"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
