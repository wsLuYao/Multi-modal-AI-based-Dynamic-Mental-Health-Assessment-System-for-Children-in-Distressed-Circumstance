"""Build an untracked DOCX from the repository's fully synthetic text sample."""

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "examples" / "synthetic_case.txt"
OUTPUT = ROOT / "runtime" / "synthetic_case.docx"


def main() -> None:
    document = Document()
    document.add_heading("合成演示个案 / Synthetic demo case", level=1)
    for paragraph in SOURCE.read_text(encoding="utf-8").split("\n\n"):
        document.add_paragraph(paragraph)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
